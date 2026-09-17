import openpyxl, json, hashlib, pathlib, datetime, math, collections

# Необязательный инструмент подготовки именно структуры приложенного Excel.
# Требуется Python с openpyxl; обычному запуску приложения Python не нужен.
import argparse
parser = argparse.ArgumentParser()
parser.add_argument('xlsx', type=pathlib.Path)
parser.add_argument('--output', type=pathlib.Path, default=pathlib.Path(__file__).resolve().parents[1] / 'data')
args = parser.parse_args()
source = args.xlsx
dest = args.output
dest.mkdir(exist_ok=True)
w = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
issues = []

def text(v): return '' if v is None else str(v).strip()
def number(v):
    if isinstance(v, (int,float)) and not isinstance(v,bool) and math.isfinite(v): return float(v)
    return None
def date(v):
    if isinstance(v, (datetime.datetime,datetime.date)): return v.strftime('%Y-%m-%d')
    if isinstance(v,str):
        for fmt in ['%d.%m.%Y','%Y-%m-%d']:
            try: return datetime.datetime.strptime(v.strip(),fmt).strftime('%Y-%m-%d')
            except ValueError: pass
    return None
def key(sheet,row): return f'{sheet}:{row}'
def warn(sheet,row,message): issues.append({'sheet':sheet,'row':row,'message':message})
def fields(sheet,row,count):
    return [w[sheet].cell(row,c).value for c in range(1,count+1)]

operations=[]
for r in range(2,1001):
    a=fields('Операции',r,7)
    if not any(v is not None for v in a): continue
    d=date(a[0]); amount=number(a[2]); typ={'Расход':'expense','Доход':'income'}.get(text(a[1]))
    if not d or amount is None or amount<=0 or not typ:
        warn('Операции',r,'Строка не перенесена: отсутствует корректная дата, тип или положительная сумма.'); continue
    operations.append({'source':key('Операции',r),'date':d,'type':typ,'amount':amount,'category':text(a[3]) or 'Без категории','subcategory':text(a[4]),'counterparty':text(a[5]),'note':text(a[6]),'title':(text(a[6]) or text(a[5]) or text(a[4]) or text(a[3]))[:100]})

products=[]
for r in range(2,65):
    a=[w['Справочники'].cell(r,c).value for c in range(7,13)]
    if not a[0]: continue
    n=[number(v) for v in a[1:]]
    if any(v is None for v in n): warn('Справочники',r,'Неполные нутриенты продукта.'); continue
    products.append({'source':key('Справочники',r),'name':text(a[0]),**dict(zip(['kcal','protein','fat','carbs','fiber'],n))})
product_map={p['name']:p for p in products}
shifts=[]
for r in range(12,999):
    a=fields('Фактическая выработка',r,15)
    if not any(a[i] is not None for i in [0,1,2,3]): continue
    d=date(a[0]); pay=number(a[2]); hours=number(a[3])
    if not d or pay is None or hours is None or hours<=0:
        warn('Фактическая выработка',r,'Неполная смена: проверь дату, оплату и часы.'); continue
    shifts.append({'source':key('Фактическая выработка',r),'date':d,'name':text(a[1]) or 'Смена','pay':pay,'hours':hours,'travelOut':number(a[4]),'travelBack':number(a[5]),'cost':number(a[8]),'tax':number(a[9]),'paidDate':date(a[12]) or '', 'status':text(a[13]) or 'не указан','note':text(a[14])})

weights=[]; measurements=[]
for r in range(6,1001):
    a=fields('Вес и тело',r,13)
    if a[0] is not None or a[1] is not None:
        if date(a[0]) and number(a[1]) is not None:
            weights.append({'source':key('Вес и тело',r),'date':date(a[0]),'weight':number(a[1]),'note':text(a[5])})
        else: warn('Вес и тело',r,'Неполная запись веса.')
    if a[7] is not None or any(a[i] is not None for i in range(8,12)):
        if date(a[7]): measurements.append({'source':key('Замеры',r),'date':date(a[7]),**dict(zip(['waist','chest','biceps','thigh'],[number(v) for v in a[8:12]])),'note':text(a[12])})
        else: warn('Вес и тело',r,'Нет даты замеров тела.')

meals=[]
for r in range(7,1001):
    a=fields('Дневник питания',r,11)
    if not any(a[i] is not None for i in range(5)): continue
    d=date(a[0]); name=text(a[3]); grams=number(a[4]); product=product_map.get(name)
    if not d or not name or grams is None or grams<=0 or not product:
        warn('Дневник питания',r,'Неполная запись питания или продукт отсутствует в справочнике.'); continue
    meals.append({'source':key('Дневник питания',r),'date':d,'time':a[1].strftime('%H:%M') if isinstance(a[1],datetime.time) else text(a[1]),'meal':text(a[2]),'name':name,'grams':grams,**{n:product[n] for n in ['kcal','protein','fat','carbs','fiber']},'note':text(a[10])})

workouts=[]
for r in range(13,1001):
    a=fields('Тренировки',r,17)
    if a[0] is None and a[1] is None: continue
    sets=[number(v) for v in a[2:10] if number(v) is not None]
    if not date(a[0]) or not a[1] or not sets:
        warn('Тренировки',r,'Неполная тренировка.'); continue
    workouts.append({'source':key('Тренировки',r),'date':date(a[0]),'name':text(a[1]),'sets':sets,'rest':number(a[10]),'note':text(a[16])})

budgets=[]
for r in range(2,15):
    name=w['Бюджет'].cell(r,1).value; amount=number(w['Бюджет'].cell(r,2).value)
    if name and amount is not None and amount>0: budgets.append({'category':text(name),'amount':amount})
goals={name:cached['Дневник питания'].cell(3,col).value for name,col in [('kcal',2),('protein',3),('fat',4),('carbs',5),('fiber',6)]}
payload={'version':1,'filename':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'operations':operations,'records':{'products':products,'shifts':shifts,'weights':weights,'measurements':measurements,'meals':meals,'workouts':workouts},'budgetTemplate':budgets,'goals':goals,'categories':[text(w['Справочники'].cell(r,2).value) for r in range(2,65) if w['Справочники'].cell(r,2).value], 'issues':issues}
summary=collections.defaultdict(lambda:{'income':0,'expense':0,'count':0})
for o in operations:
    b=summary[o['date'][:7]]; b[o['type']]+=round(o['amount']*100); b['count']+=1
payload['summary']={m:{**v,'income':v['income']/100,'expense':v['expense']/100} for m,v in summary.items()}
(dest/'excel-import.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'counts':{'operations':len(operations),**{k:len(v) for k,v in payload['records'].items()},'budgetTemplate':len(budgets)},'summary':payload['summary'],'issues':issues},ensure_ascii=False,indent=2))
